"""
Génération des exports Excel/PDF du registre phyto et de l'inventaire.
Remplace previousServer/export_clean_new.php, export_pdf.php, export_inventaire.php
(PhpSpreadsheet/TCPDF) en s'appuyant sur le nouveau schéma normalisé (au lieu des
fichiers JSON imbriqués de l'ancien système, avec leur logique de "unwrap" legacy).

Différence assumée avec l'ancien export : les traitements phyto et les travaux aux
champs sont sur deux feuilles/tables séparées plutôt que fusionnés dans une seule
feuille "cahier des charges" — les deux registres sont déjà distincts dans le nouveau
schéma (phytoTreatments / fieldWorks), les mélanger artificiellement n'apportait rien.

Limites connues (à combler si besoin) : le "Stade phénologique" (BBCH) n'est plus
capturé par le nouveau schéma ; l'unité de `treatmentTotalQuantity` et de
`surfaceTreated` n'est pas non plus stockée (seule la valeur standardisée l'est) —
ces cellules sont laissées vides plutôt que de deviner une unité.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from reportlab.lib.pagesizes import A3, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet

from app.db import cursor_appifield
from app.utils.units import find_unit_name_by_id
from app.utils.inventory_stock import compute_stock_by_product

REGISTRE_HEADERS = [
    "Année", "N° Phytolicence du responsable de l'exploitation",
    "N° d'exploitant du responsable de l'exploitation", "Date",
    "Nom de la culture ou de l'espèce emblavée", "Culture Phytoweb", "Stade phénologique",
    "Date de semis", "Remarque (culture)", "Culture archivée", "Nom de l'archive",
    "Parcelle", "Adresse", "Coordonnées", "Traitement localisé", "Type de traitement",
    "Méthode d'application", "Produit utilisé (nom commercial)", "Produit utilisé (numéro d'autorisation)", "Cible",
    "Surface traitée", "Surface traitée (unité)", "Quantité totale de traitement", "Unité",
    "Quantité de produit", "Unité", "Dose de PPP", "PPP (unité)", "Remarques",
    "Nettoyage externe de la cuve", "Nom de l'opérateur(trice)", "Heure de traitement",
]

FIELD_WORK_HEADERS = [
    "Année", "Date", "Nom de la culture ou de l'espèce emblavée", "Culture Phytoweb",
    "Date de semis", "Remarque (culture)", "Culture archivée", "Nom de l'archive",
    "Parcelle", "Type de travail", "Informations", "Outil", "Dose", "Unité", "Remarques",
]


def _year_of(value):
    if value is None:
        return ""
    if hasattr(value, "year"):
        return value.year
    return str(value)[:4]


def _fetch_user_profile(user_id: str) -> dict:
    cur = cursor_appifield()
    cur.execute("SELECT phytolicence, exploitingNumber FROM userProfile WHERE userId = %s", (user_id,))
    return cur.fetchone() or {}


def _fetch_treatment_rows(user_id: str) -> list[dict]:
    cur = cursor_appifield()
    cur.execute(
        """SELECT
            t.eventDate, t.eventTime, t.eventType, t.surfaceTreated, t.treatmentTotalQuantity,
            t.remarks AS treatmentRemarks, t.cleanFlag, t.localizedTreatment, t.treatmentType,
            pc.parcelName, pc.cropId, pc.cropName, pc.phytowebCropName,
            comp.productName, comp.productNumAut, comp.productDose, comp.productDoseUnitId,
            comp.productQuantity, comp.productQuantityUnitId, comp.productTarget,
            c.sowingDate, c.remarks AS cropRemarks, c.isArchived, c.collection,
            p.parcelAddress, p.latitude, p.longitude,
            op.operatorName, op.operatorPhytolicence
           FROM phytoTreatments t
           JOIN phytoTreatmentsParcelsCrops pc ON pc.eventId = t.eventId
           JOIN phytoTreatmentsComposition comp ON comp.eventId = t.eventId
           LEFT JOIN crops c ON c.cropId = pc.cropId AND c.userId = t.userId
           LEFT JOIN parcels p ON p.parcelId = pc.parcelId AND p.userId = t.userId
           LEFT JOIN operators op ON op.operatorId = t.operatorId AND op.userId = t.userId
           WHERE t.userId = %s
           ORDER BY t.eventDate""",
        (user_id,),
    )
    return cur.fetchall()


def _fetch_field_work_rows(user_id: str) -> list[dict]:
    cur = cursor_appifield()
    cur.execute(
        """SELECT
            fw.eventDate, fw.eventType, fw.eventInfo, fw.remarks, fw.eventDose,
            fw.eventDoseUnitId, fw.eventTool,
            fwpc.parcelName, fwpc.cropId, fwpc.cropName, fwpc.phytowebCropName,
            c.sowingDate, c.remarks AS cropRemarks, c.isArchived, c.collection
           FROM fieldWorks fw
           JOIN fieldWorksParcelsCrops fwpc ON fwpc.eventId = fw.eventId
           LEFT JOIN crops c ON c.cropId = fwpc.cropId AND c.userId = fw.userId
           WHERE fw.userId = %s
           ORDER BY fw.eventDate""",
        (user_id,),
    )
    return cur.fetchall()


def _apply_filters(rows: list[dict], filters: dict) -> list[dict]:
    years = set(filters.get("years") or [])
    crop_names = set(filters.get("cropNames") or [])
    parcels = set(filters.get("parcels") or [])
    archive_status = filters.get("archiveStatus", "all")

    def keep(row):
        if years and str(_year_of(row.get("eventDate"))) not in years:
            return False
        if crop_names and row.get("cropName") not in crop_names:
            return False
        if parcels and row.get("parcelName") not in parcels:
            return False
        if archive_status == "archived" and not row.get("isArchived"):
            return False
        if archive_status == "not-archived" and row.get("isArchived"):
            return False
        return True

    return [row for row in rows if keep(row)]


def _operator_label(row: dict) -> str:
    name = row.get("operatorName")
    if not name:
        return ""
    licence = row.get("operatorPhytolicence")
    return f"{name} ({licence})" if licence else name


def _registre_row_values(row: dict, profile: dict) -> list:
    return [
        _year_of(row.get("eventDate")),
        profile.get("phytolicence") or "",
        profile.get("exploitingNumber") or "",
        row.get("eventDate"),
        row.get("cropName") or "",
        row.get("phytowebCropName") or "",
        "",  # Stade phénologique (BBCH) : non modélisé dans le nouveau schéma
        row.get("sowingDate") or "",
        row.get("cropRemarks") or "",
        "Vrai" if row.get("isArchived") else "Faux",
        row.get("collection") if row.get("isArchived") else "",
        row.get("parcelName") or "",
        row.get("parcelAddress") or "",
        f"{row['latitude']}, {row['longitude']}" if row.get("latitude") is not None else "",
        "Oui" if row.get("localizedTreatment") else "Non",
        row.get("eventType") or "",
        (row.get("treatmentType") or "").capitalize(),  # stocké en minuscules (ex. "pulvérisation")
        row.get("productName") or "",
        row.get("productNumAut") or "",
        row.get("productTarget") or "",
        row.get("surfaceTreated"),
        "",  # unité de surfaceTreated non stockée
        row.get("treatmentTotalQuantity"),
        "",  # unité de treatmentTotalQuantity non stockée
        row.get("productQuantity"),
        find_unit_name_by_id(row.get("productQuantityUnitId")),
        row.get("productDose"),
        find_unit_name_by_id(row.get("productDoseUnitId")),
        row.get("treatmentRemarks") or "",
        "Oui" if row.get("cleanFlag") else "Non",
        _operator_label(row),
        str(row.get("eventTime") or ""),
    ]


def _field_work_row_values(row: dict) -> list:
    return [
        _year_of(row.get("eventDate")),
        row.get("eventDate"),
        row.get("cropName") or "",
        row.get("phytowebCropName") or "",
        row.get("sowingDate") or "",
        row.get("cropRemarks") or "",
        "Vrai" if row.get("isArchived") else "Faux",
        row.get("collection") if row.get("isArchived") else "",
        row.get("parcelName") or "",
        row.get("eventType") or "",
        row.get("eventInfo") or "",
        row.get("eventTool") or "",
        row.get("eventDose"),
        find_unit_name_by_id(row.get("eventDoseUnitId")),
        row.get("remarks") or "",
    ]


def _write_sheet(sheet, headers: list[str], data_rows: list[list]):
    for col_idx, title in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=title)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row_idx, values in enumerate(data_rows, start=2):
        for col_idx, value in enumerate(values, start=1):
            sheet.cell(row=row_idx, column=col_idx, value=value)


def _registre_data(user_id: str, filters: dict):
    profile = _fetch_user_profile(user_id)
    treatment_rows = _apply_filters(_fetch_treatment_rows(user_id), filters)
    field_work_rows = _apply_filters(_fetch_field_work_rows(user_id), filters)
    return (
        [_registre_row_values(row, profile) for row in treatment_rows],
        [_field_work_row_values(row) for row in field_work_rows],
    )


def build_registre_excel(user_id: str, filters: dict) -> BytesIO:
    treatment_data, field_work_data = _registre_data(user_id, filters)

    wb = Workbook()
    _write_sheet(wb.active, REGISTRE_HEADERS, treatment_data)
    wb.active.title = "Registre PPP"
    _write_sheet(wb.create_sheet("Travaux aux champs"), FIELD_WORK_HEADERS, field_work_data)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _pdf_table(headers, data_rows):
    table_data = [headers] + [["" if v is None else str(v) for v in row] for row in data_rows]
    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2c5f2d")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 6),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    return table


def build_registre_pdf(user_id: str, filters: dict) -> BytesIO:
    treatment_data, field_work_data = _registre_data(user_id, filters)

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A3))
    styles = getSampleStyleSheet()
    elements = [
        Paragraph("Registre PPP", styles["Heading1"]),
        Spacer(1, 12),
        _pdf_table(REGISTRE_HEADERS, treatment_data),
        PageBreak(),
        Paragraph("Travaux aux champs", styles["Heading1"]),
        Spacer(1, 12),
        _pdf_table(FIELD_WORK_HEADERS, field_work_data),
    ]
    doc.build(elements)
    buffer.seek(0)
    return buffer


def build_inventaire_excel(user_id: str) -> BytesIO:
    stock = compute_stock_by_product(user_id)

    wb = Workbook()
    summary_headers = ["Produit (id)", "Stock disponible", "Stock PPNU", "Stock PPNV",
                       "Corrections (net)", "Nombre de mouvements"]
    summary_rows = [
        [entry["productId"], entry["stock"], entry["ppnuStock"], entry["ppnvStock"],
         entry["correctionStock"], len(entry["operations"])]
        for entry in stock
    ]
    _write_sheet(wb.active, summary_headers, summary_rows)
    wb.active.title = "Inventaire"

    detail_headers = ["Produit (id)", "Type", "Date", "Quantité", "Unité", "Traitement source"]
    detail_rows = [
        [
            entry["productId"],
            # Le sens d'une correction n'est pas porté par la quantité (toujours > 0).
            f"correction ({op['correctionMode']})" if op["operationType"] == "correction" else op["operationType"],
            op["operationDate"],
            op.get("quantity"),
            find_unit_name_by_id(op.get("quantityUnitId")),
            op.get("sourceEventId") or "",
        ]
        for entry in stock
        for op in entry["operations"]
    ]
    _write_sheet(wb.create_sheet("Mouvements"), detail_headers, detail_rows)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
