"""
Import Excel d'achats de produits phyto dans l'inventaire (onglet Inventaire > Importer des achats).

Trois étapes, aucune écriture avant la dernière :
  1. build_template()          : classeur vierge à compléter (GET /inventories/import/template)
  2. preview_import(...)       : lecture + validation ligne par ligne, avec conflits et
                                 suggestions de produits (POST /inventories/import/preview)
  3. validate_confirm_rows(...): revalidation des lignes corrigées par l'utilisateur avant
                                 insertion (POST /inventories/import/confirm)

Le produit est identifié par son NUMÉRO D'AUTORISATION (general.numAut, actes les plus
récents) ; la colonne « Nom du produit » du fichier est ignorée pour l'import (le nom vient
de Phytoweb) et ne sert qu'à proposer des suggestions quand le numéro n'est pas reconnu.
"""
import re
import unicodedata
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.db import cursor_appifield, cursor_phytoweb

TEMPLATE_HEADERS = ["Numéro d'autorisation", "Nom du produit", "Quantité", "Unité", "Date d'achat"]
_FIELDS = ("numAut", "productName", "quantity", "unit", "date")
MAX_ROWS = 500
MAX_FILE_SIZE = 2 * 1024 * 1024
ALLOWED_UNITS = ("L", "Kg")  # noms d'unité attendus par POST /inventories (cf. achat-de-ppp)

_UNIT_ALIASES = {
    "l": "L", "litre": "L", "litres": "L", "liter": "L", "liters": "L", "lt": "L", "ltr": "L",
    "kg": "Kg", "kgs": "Kg", "kilo": "Kg", "kilos": "Kg", "kilogramme": "Kg", "kilogrammes": "Kg",
}


class ImportFileError(Exception):
    """Fichier illisible ou au mauvais format (erreur globale, pas une erreur de ligne)."""


# ── Template ──────────────────────────────────────────────────────────────────

def build_template() -> BytesIO:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Achats"
    header_fill = PatternFill("solid", fgColor="E3F6F3")
    for col_idx, title in enumerate(TEMPLATE_HEADERS, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=title)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for col, width in zip("ABCDE", (24, 32, 12, 10, 16)):
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = "A2"

    last = MAX_ROWS + 1
    units = DataValidation(type="list", formula1='"L,kg"', allow_blank=True,
                           error="Unité attendue : L ou kg", errorTitle="Unité")
    sheet.add_data_validation(units)
    units.add(f"D2:D{last}")
    for row in range(2, last + 1):
        sheet.cell(row=row, column=1).number_format = "@"          # 10992P/B : texte
        sheet.cell(row=row, column=5).number_format = "DD/MM/YYYY"

    help_sheet = wb.create_sheet("Instructions")
    lines = [
        ("Import d'achats de produits phytopharmaceutiques", True),
        ("Complétez l'onglet « Achats », une ligne par achat, puis importez le fichier dans Appi'Field.", False),
        ("", False),
        ("Numéro d'autorisation : obligatoire, tel qu'indiqué sur l'étiquette (ex. 10992P/B).", False),
        ("Nom du produit : facultatif, pour votre repère uniquement (le nom officiel Phytoweb est utilisé).", False),
        ("Quantité : obligatoire, nombre supérieur à 0 (ex. 12,5).", False),
        ("Unité : obligatoire, L ou kg.", False),
        ("Date d'achat : obligatoire, au format JJ/MM/AAAA (ex. 15/03/2026).", False),
        ("", False),
        ("Exemple :", True),
        ("10992P/B | TUSK | 20 | L | 15/03/2026", False),
    ]
    for idx, (text, bold) in enumerate(lines, start=1):
        help_sheet.cell(row=idx, column=1, value=text).font = Font(bold=bold)
    help_sheet.column_dimensions["A"].width = 100

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ── Normalisation des valeurs ─────────────────────────────────────────────────

def _strip_accents(value: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", value) if unicodedata.category(c) != "Mn")


def _header_key(value) -> str:
    return re.sub(r"[^a-z]", "", _strip_accents(str(value or "")).lower())


_HEADER_TO_FIELD = {
    "numerodautorisation": "numAut", "numeroautorisation": "numAut", "numaut": "numAut",
    "nautorisation": "numAut", "autorisation": "numAut",
    "nomduproduit": "productName", "nomproduit": "productName", "produit": "productName",
    "quantite": "quantity", "qte": "quantity",
    "unite": "unit",
    "datedachat": "date", "dateachat": "date", "date": "date",
}


def normalize_num_aut(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", "", str(value)).upper()


def _display(value) -> str:
    """Valeur brute du fichier, telle qu'à afficher dans la popup de résolution."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_quantity(value):
    """Retourne (float, None) ou (None, message d'erreur)."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, "Quantité manquante"
    if isinstance(value, bool):
        return None, "Quantité invalide"
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = str(value).strip().replace(" ", "").replace(" ", "").replace(",", ".")
        try:
            number = float(text)
        except ValueError:
            return None, f"Quantité invalide : « {value} »"
    if number <= 0:
        return None, "La quantité doit être supérieure à 0"
    return number, None


def parse_unit(value):
    """Retourne ('L'|'Kg', None) ou (None, message d'erreur)."""
    text = _strip_accents(str(value or "")).strip().lower().rstrip(".")
    if not text:
        return None, "Unité manquante (L ou kg)"
    unit = _UNIT_ALIASES.get(text)
    if not unit:
        return None, f"Unité non reconnue : « {value} » (L ou kg attendu)"
    return unit, None


def parse_purchase_date(value):
    """Cellule date Excel ou texte JJ/MM/AAAA -> ('YYYY-MM-DD', None) ou (None, message)."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, "Date d'achat manquante"
    if isinstance(value, datetime):
        parsed = value.date()
    elif isinstance(value, date):
        parsed = value
    else:
        match = re.fullmatch(r"\s*(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{4})\s*", str(value))
        if not match:
            return None, f"Date invalide : « {value} » (format JJ/MM/AAAA attendu)"
        day, month, year = (int(g) for g in match.groups())
        try:
            parsed = date(year, month, day)
        except ValueError:
            return None, f"Date inexistante : « {value} »"
    if parsed > date.today():
        return None, "La date d'achat ne peut pas être dans le futur"
    return parsed.isoformat(), None


def parse_iso_date(value):
    """Date déjà normalisée renvoyée par le frontend à l'étape confirm (YYYY-MM-DD)."""
    try:
        parsed = date.fromisoformat(str(value))
    except (TypeError, ValueError):
        return None, "operationDate invalide (YYYY-MM-DD attendu)"
    if parsed > date.today():
        return None, "La date d'achat ne peut pas être dans le futur"
    return parsed.isoformat(), None


# ── Lecture du fichier ────────────────────────────────────────────────────────

def read_rows(file_bytes: bytes) -> list[dict]:
    """Lignes non vides de la feuille « Achats » (ou de la première feuille), avec leur n° de ligne Excel."""
    try:
        wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFileError("Fichier Excel illisible (format .xlsx attendu)") from exc

    sheet = wb["Achats"] if "Achats" in wb.sheetnames else wb.worksheets[0]
    rows_iter = sheet.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        raise ImportFileError("Le fichier est vide")

    columns = {}
    for idx, title in enumerate(header):
        field = _HEADER_TO_FIELD.get(_header_key(title))
        if field and field not in columns:
            columns[field] = idx
    missing = [TEMPLATE_HEADERS[_FIELDS.index(f)] for f in ("numAut", "quantity", "unit", "date") if f not in columns]
    if missing:
        raise ImportFileError(
            "Colonnes introuvables : " + ", ".join(missing) + ". Utilisez le modèle d'import."
        )

    rows = []
    for excel_row, values in enumerate(rows_iter, start=2):
        record = {f: (values[i] if i < len(values) else None) for f, i in columns.items()}
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in record.values()):
            continue
        record["rowNumber"] = excel_row
        rows.append(record)
        if len(rows) > MAX_ROWS:
            raise ImportFileError(f"Trop de lignes : {MAX_ROWS} achats maximum par import")
    if not rows:
        raise ImportFileError("Aucun achat trouvé dans le fichier")
    return rows


# ── Référentiel produits ──────────────────────────────────────────────────────

def _load_recent_products() -> dict[str, dict]:
    """{numAut normalisé: {idAct, numAut, productName, statut}} pour les actes les plus récents."""
    cur = cursor_phytoweb()
    cur.execute("SELECT idAct, numAut, productName, statut FROM general WHERE mostRecentAct = 1")
    return {normalize_num_aut(row["numAut"]): row for row in cur.fetchall()}


def _product_payload(row: dict) -> dict:
    return {"productId": row["idAct"], "numAut": row["numAut"], "productName": row["productName"]}


def _suggestions(num_aut: str, product_name: str, products: dict[str, dict], limit: int = 5) -> list[dict]:
    """Produits proches pour un numéro non reconnu : même numéro sans/avec autre suffixe, puis nom."""
    found: list[dict] = []
    seen = set()

    def add(row):
        if row["idAct"] not in seen and len(found) < limit:
            seen.add(row["idAct"])
            found.append(_product_payload(row))

    digits = re.match(r"\d+", num_aut or "")
    if digits:
        for key, row in products.items():
            if re.match(r"\d+", key) and re.match(r"\d+", key).group() == digits.group():
                add(row)
    name_key = _strip_accents((product_name or "").strip().lower())
    if len(name_key) >= 3:
        for row in products.values():
            if name_key in _strip_accents((row["productName"] or "").lower()):
                add(row)
    return found


def _existing_purchase_keys(user_id: str) -> set[tuple]:
    cur = cursor_appifield()
    cur.execute(
        """SELECT productId, operationDate, quantity FROM inventoriesOperations
           WHERE userId = %s AND operationType = 'purchase'""",
        (user_id,),
    )
    return {
        (row["productId"], str(row["operationDate"]), round(float(row["quantity"] or 0), 6))
        for row in cur.fetchall()
    }


# ── Étape 2 : prévisualisation ───────────────────────────────────────────────

def preview_import(user_id: str, file_bytes: bytes) -> dict:
    raw_rows = read_rows(file_bytes)
    products = _load_recent_products()
    existing = _existing_purchase_keys(user_id)
    seen_in_file: dict[tuple, int] = {}

    result_rows = []
    for raw in raw_rows:
        errors, warnings = [], []
        num_aut_input = _display(raw.get("numAut"))
        file_name = _display(raw.get("productName"))

        product = None
        suggestions = []
        if not num_aut_input:
            errors.append({"field": "numAut", "message": "Numéro d'autorisation manquant"})
            suggestions = _suggestions("", file_name, products)
        else:
            product = products.get(normalize_num_aut(num_aut_input))
            if not product:
                errors.append({"field": "numAut",
                               "message": f"Numéro d'autorisation non reconnu : « {num_aut_input} »"})
                suggestions = _suggestions(normalize_num_aut(num_aut_input), file_name, products)
            elif (product.get("statut") or "").strip() == "Retrait":
                warnings.append({"field": "numAut", "message": "Ce produit n'est plus autorisé (statut : Retrait)"})

        quantity, error = parse_quantity(raw.get("quantity"))
        if error:
            errors.append({"field": "quantity", "message": error})
        unit, error = parse_unit(raw.get("unit"))
        if error:
            errors.append({"field": "unit", "message": error})
        operation_date, error = parse_purchase_date(raw.get("date"))
        if error:
            errors.append({"field": "date", "message": error})

        if product and quantity is not None and operation_date:
            key = (product["idAct"], operation_date, round(quantity, 6))
            if key in existing:
                warnings.append({"field": "row", "message": "Un achat identique (produit, date, quantité) existe déjà dans votre inventaire"})
            if key in seen_in_file:
                warnings.append({"field": "row", "message": f"Doublon de la ligne {seen_in_file[key]} du fichier"})
            else:
                seen_in_file[key] = raw["rowNumber"]

        result_rows.append({
            "rowNumber": raw["rowNumber"],
            "input": {
                "numAut": num_aut_input,
                "productName": file_name,
                "quantity": _display(raw.get("quantity")),
                "unit": _display(raw.get("unit")),
                "date": _display(raw.get("date")),
            },
            "product": _product_payload(product) if product else None,
            "quantity": quantity,
            "quantityUnit": unit,
            "operationDate": operation_date,
            "errors": errors,
            "warnings": warnings,
            "suggestions": suggestions,
            "status": "error" if errors else ("warning" if warnings else "ok"),
        })

    return {
        "rows": result_rows,
        "summary": {
            "total": len(result_rows),
            "ok": sum(1 for r in result_rows if r["status"] == "ok"),
            "warning": sum(1 for r in result_rows if r["status"] == "warning"),
            "error": sum(1 for r in result_rows if r["status"] == "error"),
        },
    }


# ── Étape 3 : revalidation avant insertion ───────────────────────────────────

def validate_confirm_rows(rows) -> tuple[list[dict], list[dict]]:
    """
    Revalide les lignes envoyées à /import/confirm (déjà corrigées dans la popup).
    Retourne (lignes valides normalisées, erreurs [{index, rowNumber, field, message}]).
    """
    if not isinstance(rows, list) or not rows:
        return [], [{"index": None, "rowNumber": None, "field": "rows", "message": "Aucun achat à importer"}]
    if len(rows) > MAX_ROWS:
        return [], [{"index": None, "rowNumber": None, "field": "rows",
                     "message": f"{MAX_ROWS} achats maximum par import"}]

    product_ids = {str(r.get("productId")) for r in rows if isinstance(r, dict) and r.get("productId")}
    known_ids = set()
    if product_ids:
        cur = cursor_phytoweb()
        placeholders = ", ".join(["%s"] * len(product_ids))
        cur.execute(f"SELECT idAct FROM general WHERE idAct IN ({placeholders})", tuple(product_ids))
        known_ids = {row["idAct"] for row in cur.fetchall()}

    valid, errors = [], []
    for index, row in enumerate(rows):
        row = row if isinstance(row, dict) else {}
        row_errors = []

        def fail(field, message):
            row_errors.append({"index": index, "rowNumber": row.get("rowNumber"), "field": field, "message": message})

        product_id = str(row.get("productId") or "")
        if product_id not in known_ids:
            fail("productId", "Produit Phytoweb inconnu")
        quantity, error = parse_quantity(row.get("quantity"))
        if error:
            fail("quantity", error)
        unit = row.get("quantityUnit")
        if unit not in ALLOWED_UNITS:
            fail("quantityUnit", "Unité attendue : L ou Kg")
        operation_date, error = parse_iso_date(row.get("operationDate"))
        if error:
            fail("operationDate", error)

        if row_errors:
            errors.extend(row_errors)
        else:
            valid.append({"productId": product_id, "quantity": quantity,
                          "quantityUnit": unit, "operationDate": operation_date})
    return valid, errors
